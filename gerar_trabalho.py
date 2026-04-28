"""
Gera o trabalho acadêmico Pulsar RH em formato DOCX com normas ABNT.
Execução: python gerar_trabalho.py
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = (
    r"d:\OneDrive\Área de Trabalho\_projetos-programacao"
    r"\pulsarRH\public\trabalho-pulsar-rh.docx"
)


# ──────────────────────────────────────────────────────────────
# Utilitários de formatação
# ──────────────────────────────────────────────────────────────

def _fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=Cm(1.25),
         left_indent=Cm(0), sb=0, sa=0,
         spacing=WD_LINE_SPACING.ONE_POINT_FIVE):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = first_indent
    pf.left_indent = left_indent
    pf.space_before = Pt(sb)
    pf.space_after = Pt(sa)
    pf.line_spacing_rule = spacing


def _run(p, text, bold=False, size=12, font="Times New Roman"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    return r


def body(doc, text):
    p = doc.add_paragraph()
    _fmt(p)
    _run(p, text)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    _fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0), sb=24, sa=6)
    _run(p, text.upper(), bold=True)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    _fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=Cm(0), sb=12, sa=6)
    _run(p, text, bold=True)
    return p


def capa_linha(doc, text, bold=False, size=12, sb=0, sa=0):
    p = doc.add_paragraph()
    _fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0), sb=sb, sa=sa)
    _run(p, text, bold=bold, size=size)
    return p


def espacador(doc, qtd=1):
    for _ in range(qtd):
        p = doc.add_paragraph()
        _fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=Cm(0))
        _run(p, "")


def code(doc, text):
    """Bloco de código: cada linha como parágrafo com fonte monoespaçada."""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(1.25)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(6) if i == 0 else Pt(0)
        pf.space_after = Pt(6) if i == len(lines) - 1 else Pt(0)
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(10)


def referencia(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(12)
    _run(p, text)
    return p


def add_page_number(section):
    header = section.header
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.first_line_indent = Cm(0)
    r = hp.add_run()
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    for tag, ftype in [("begin", None), (None, "PAGE"), ("end", None)]:
        if ftype:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = ftype
            r._r.append(el)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            r._r.append(fc)


# ──────────────────────────────────────────────────────────────
# Capa
# ──────────────────────────────────────────────────────────────

def build_capa(doc):
    capa_linha(doc, "UNIFECAF – FACULDADE DE TECNOLOGIA ROCKETSEAT", bold=True)
    capa_linha(doc, "Curso de Inteligência Artificial e Automação Digital")
    espacador(doc, 10)
    capa_linha(doc, "YURE DOS SANTOS SANTANA", bold=True, sa=48)
    capa_linha(
        doc,
        "SEU PRIMEIRO COPILOTO DE IA: CRIANDO UMA SOLUÇÃO INTELIGENTE "
        "COM IA GENERATIVA",
        bold=True, sa=24,
    )
    capa_linha(
        doc,
        "Trabalho apresentado à disciplina de Fundamentos de I.A. Generativa "
        "do curso de Inteligência Artificial e Automação Digital da UniFeCAF – "
        "Faculdade de Tecnologia Rocketseat, como requisito de avaliação.",
    )
    espacador(doc, 8)
    capa_linha(doc, "SALVADOR")
    capa_linha(doc, "2025")


# ──────────────────────────────────────────────────────────────
# Sumário
# ──────────────────────────────────────────────────────────────

def build_sumario(doc):
    h1(doc, "SUMÁRIO")
    entries = [
        ("1", "CONTEXTUALIZAÇÃO DO DESAFIO", False),
        ("2", "JUSTIFICATIVA PARA USO DE IA GENERATIVA", False),
        ("3", "MODELO LLM UTILIZADO: GOOGLE GEMINI 2.5 FLASH", False),
        ("4", "PROMPT ENGINEERING: ELABORAÇÃO DO PROMPT", False),
        ("5", "DESCRIÇÃO DA SOLUÇÃO", False),
        ("5.1", "Tecnologias utilizadas", True),
        ("5.2", "Fluxo de funcionamento", True),
        ("5.3", "Arquitetura da automação no Make.com", True),
        ("6", "BENEFÍCIOS PERCEBIDOS E DESAFIOS ENFRENTADOS", False),
        ("7", "LIMITES ÉTICOS E DE SEGURANÇA", False),
        ("7.1", "Privacidade de dados e LGPD", True),
        ("7.2", "Viés da IA", True),
        ("7.3", "Segurança das credenciais", True),
        ("", "REFERÊNCIAS", False),
    ]
    for num, titulo, is_sub in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0.5) if is_sub else Cm(0)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        label = f"{num}  {titulo}" if num else titulo
        _run(p, label, bold=not is_sub)


# ──────────────────────────────────────────────────────────────
# Seções do corpo
# ──────────────────────────────────────────────────────────────

def build_s1(doc):
    h1(doc, "1 CONTEXTUALIZAÇÃO DO DESAFIO")
    body(doc,
         "O setor de Recursos Humanos em empresas brasileiras lida diariamente com um "
         "volume expressivo de demandas de comunicação escrita repetitiva: e-mails de "
         "convocação, resumos de reuniões, mensagens para grupos corporativos no WhatsApp "
         "e avisos institucionais. Essas atividades, embora necessárias, consomem tempo de "
         "profissionais que poderiam estar dedicados a ações mais estratégicas, como "
         "desenvolvimento de talentos, gestão de clima organizacional e processos seletivos.")
    body(doc,
         "O desafio proposto pela disciplina de Fundamentos de I.A. Generativa consiste em "
         "idealizar e prototipar uma solução baseada em Modelos de Linguagem de Grande Escala "
         "(LLMs) com engenharia de prompts para automatizar esse tipo de tarefa, sem exigir "
         "programação complexa do usuário final. O objetivo é criar um assistente inteligente "
         "capaz de gerar mensagens corporativas coerentes, bem escritas e com identidade "
         "organizacional a partir de entradas simples.")
    body(doc,
         "O projeto Pulsar RH foi desenvolvido como resposta direta a esse desafio: uma "
         "aplicação web que permite que colaboradores e equipes de RH gerem mensagens "
         "corporativas completas e prontas para uso, a partir do preenchimento de apenas "
         "cinco campos contextuais. A ferramenta está disponível publicamente em: "
         "https://santanayure.github.io/pulsarRH/.")


def build_s2(doc):
    h1(doc, "2 JUSTIFICATIVA PARA USO DE IA GENERATIVA")
    body(doc,
         "Os Modelos de Linguagem de Grande Escala (LLMs) são especialmente adequados para "
         "tarefas de geração de texto porque foram treinados sobre corpora de texto massivos "
         "e desenvolveram a capacidade de compreender e produzir linguagem natural em múltiplos "
         "estilos, registros e formatos. Diferentemente de sistemas baseados em regras ou "
         "templates fixos, os LLMs se adaptam às variações sutis de contexto, mantendo "
         "coerência semântica e adequação formal.")
    body(doc,
         "A aplicação de IA Generativa ao desafio proposto justifica-se por quatro vantagens "
         "centrais. A primeira é a velocidade: o tempo de elaboração de uma comunicação passa "
         "de minutos de edição manual para segundos de geração automática. A segunda é a "
         "consistência: a ferramenta mantém a identidade comunicacional da organização "
         "independentemente de quem está fazendo a solicitação. A terceira é a acessibilidade: "
         "colaboradores com menos habilidade em escrita formal recebem suporte de qualidade "
         "equivalente ao de profissionais experientes. A quarta é a escalabilidade: a mesma "
         "solução atende a um número ilimitado de solicitações simultâneas sem custo marginal "
         "adicional de pessoal.")
    body(doc,
         "A escolha do Make.com como camada de integração, em vez de chamar a API do Gemini "
         "diretamente a partir do navegador, foi uma decisão arquitetural deliberada: as "
         "credenciais de API nunca são expostas no código-fonte do frontend, e o fluxo de "
         "automação pode ser expandido com novos módulos — como registro em banco de dados ou "
         "envio automático de e-mail — sem necessidade de alterar a interface do usuário.")


def build_s3(doc):
    h1(doc, "3 MODELO LLM UTILIZADO: GOOGLE GEMINI 2.5 FLASH")
    body(doc,
         "O Gemini é a família de modelos de linguagem multimodal desenvolvida pelo Google "
         "DeepMind, capaz de processar e gerar texto, imagens, áudio e código. Lançada em "
         "2025, a série Gemini 2.5 representa a geração mais recente e capaz da plataforma. "
         "A variante Flash foi escolhida para este projeto com base em três critérios: "
         "latência, custo e adequação à tarefa.")
    body(doc,
         "No que se refere à latência, os modelos Flash são otimizados para velocidade de "
         "resposta, sendo essenciais em interfaces interativas onde o usuário espera retorno "
         "imediato. Em relação ao custo, a variante Flash apresenta custo por token "
         "significativamente menor em comparação com as variantes Pro, sendo mais adequada "
         "para um projeto acadêmico com orçamento limitado. Quanto à qualidade, mesmo a "
         "versão econômica do Gemini 2.5 demonstra capacidade suficiente para geração de "
         "texto profissional estruturado, uma vez que o prompt fornece contexto claro e "
         "delimitado para a tarefa.")
    body(doc,
         "O Gemini 2.5 possui suporte nativo ao português (Brasil) e foi treinado sobre "
         "corpus multilíngue diverso, incluindo textos corporativos e comunicações formais, "
         "o que contribui para a qualidade das saídas geradas. A integração foi realizada "
         "por meio do módulo nativo do Make.com para a API Gemini, que gerencia a "
         "autenticação via chave de API armazenada de forma segura no cofre da plataforma, "
         "sem exposição no código-fonte.")


def build_s4(doc):
    h1(doc, "4 PROMPT ENGINEERING: ELABORAÇÃO DO PROMPT")
    body(doc,
         "A eficácia de qualquer solução baseada em LLM depende criticamente de como o "
         "prompt é estruturado. Para o Pulsar RH, foram projetadas duas camadas de prompt: "
         "o system prompt, que define o comportamento permanente do modelo, e o prompt do "
         "usuário, construído dinamicamente com os dados do formulário.")

    h2(doc, "4.1 System Prompt")
    body(doc,
         "O system prompt define a persona do modelo e as regras de comportamento antes "
         "que ele receba os dados do usuário. Sua elaboração seguiu práticas consolidadas "
         "de engenharia de prompts (WHITE et al., 2023), em especial os padrões de "
         '"role prompting" — atribuição de uma persona especializada ao modelo — e '
         '"output constraints" — definição explícita de restrições para o formato da saída. '
         "O system prompt utilizado na automação é transcrito a seguir:")
    code(doc,
         "Você é um especialista em comunicação corporativa de RH, com foco em clareza,\n"
         "profissionalismo e padronização institucional.\n"
         "\n"
         "Sua tarefa é gerar uma mensagem corporativa com base nos dados fornecidos.\n"
         "\n"
         "Antes de escrever, interprete corretamente o contexto e o objetivo da mensagem.\n"
         "\n"
         "INSTRUÇÕES:\n"
         "1. Adapte o formato conforme o tipo de texto:\n"
         "   - E-mail: incluir saudação e fechamento\n"
         "   - WhatsApp: direto, objetivo e natural\n"
         "   - Resumo: organizado e claro\n"
         "   - Aviso institucional: formal e padronizado\n"
         "\n"
         "2. Ajuste o estilo conforme o tom de voz:\n"
         "   - Formal: linguagem profissional\n"
         "   - Informal: linguagem leve e acessível\n"
         "   - Urgente: direto e com senso de prioridade\n"
         "   - Acolhedor: empático e humanizado\n"
         "\n"
         "3. Utilize o tema como base principal da mensagem\n"
         "4. Não invente informações que não foram fornecidas\n"
         "5. Mantenha coerência, clareza e objetividade\n"
         "\n"
         "REGRAS (OBRIGATÓRIAS):\n"
         "- Escreva em português do Brasil\n"
         "- Não explique o que você está fazendo\n"
         "- Não use listas ou tópicos\n"
         "- Retorne apenas a mensagem final pronta para uso")

    h2(doc, "4.2 Prompt do Usuário")
    body(doc,
         "O prompt do usuário é construído dinamicamente pelo Make.com, mapeando cada "
         "campo do formulário para uma variável do cenário (notação {{módulo.campo}}). "
         "O template utilizado é o seguinte:")
    code(doc,
         "DADOS DE ENTRADA:\n"
         "- Nome e cargo: {{1.nomeCargo}}\n"
         "- Tema ou assunto: {{1.temaAssunto}}\n"
         "- Tipo de texto: {{1.tipoTexto}}\n"
         "- Tom de voz: {{1.tomVoz}}\n"
         "\n"
         "SAÍDA ESPERADA:\n"
         "Uma única mensagem completa, pronta para envio, adequada ao contexto\n"
         "e ao tipo de comunicação solicitado.")
    body(doc,
         "A combinação das duas camadas segue as melhores práticas de engenharia de prompts: "
         "separação de responsabilidades (regras gerais no sistema, dados na mensagem do "
         "usuário), persona especializada, restrições explícitas de saída e ancoragem "
         "das expectativas com a seção 'SAÍDA ESPERADA', que orienta o modelo quanto "
         "ao formato e extensão da resposta esperada.")


def build_s5(doc):
    h1(doc, "5 DESCRIÇÃO DA SOLUÇÃO")

    h2(doc, "5.1 Tecnologias Utilizadas")
    body(doc,
         "O frontend foi desenvolvido com React 19, TypeScript 5 e Vite 8, com CSS Modules "
         "para estilização escopada por componente. A arquitetura utiliza o padrão de hooks "
         "customizados para gerenciamento de estado (useMessageForm) e o padrão de serviço "
         "com injeção de dependência para abstrair a integração com o backend "
         "(MakeWebhookSubmissionService). A aplicação não possui servidor próprio: toda a "
         "lógica de negócio e processamento de IA é gerenciada pelo Make.com e pelo Google.")
    body(doc,
         "O Make.com atua como camada de integração sem código: recebe os dados do "
         "formulário via webhook HTTP, orquestra a chamada ao Google Gemini 2.5 Flash e "
         "retorna a mensagem gerada ao frontend em formato JSON. A implantação é "
         "automatizada via GitHub Actions e o artefato final é servido pelo GitHub Pages, "
         "sem custo de infraestrutura.")

    h2(doc, "5.2 Fluxo de Funcionamento")
    body(doc,
         "O fluxo completo da solução percorre seis etapas. Na primeira, o usuário acessa "
         "a aplicação e preenche os cinco campos do formulário: endereço de e-mail, nome e "
         "cargo, tema ou assunto, tipo de texto (E-mail, Resumo, WhatsApp ou Aviso "
         "Institucional) e tom de voz desejado (Formal, Informal, Urgente ou Acolhedor).")
    body(doc,
         "Na segunda etapa, ao submeter o formulário, o frontend serializa os dados em JSON "
         "e realiza uma requisição HTTP POST para a URL do webhook do Make.com, armazenada "
         "como variável de ambiente. Na terceira etapa, o Make.com recebe os dados e aciona "
         "o módulo do Google Gemini, compondo o prompt com o system prompt configurado e os "
         "dados mapeados do formulário.")
    body(doc,
         "Na quarta etapa, o Gemini processa o prompt e gera a mensagem corporativa. Na "
         "quinta etapa, o Make.com retorna ao frontend um objeto JSON no formato "
         '{"success": true, "mensagem": "..."}. Por fim, o frontend exibe a mensagem '
         "gerada diretamente no formulário, em destaque visual, pronta para que o usuário "
         "copie e utilize em seu canal de comunicação.")

    h2(doc, "5.3 Arquitetura da Automação no Make.com")
    body(doc,
         "O cenário no Make.com é composto por três módulos executados em sequência. O "
         "primeiro é o Custom Webhook (gateway:CustomWebHook), responsável por receber "
         "a requisição HTTP POST do frontend e disponibilizar os campos do payload JSON "
         "como variáveis para os módulos seguintes. O cenário está configurado no modo "
         "instantâneo (instant: true), respondendo imediatamente ao receber os dados.")
    body(doc,
         "O segundo módulo é o Google Gemini AI "
         "(gemini-ai:createACompletionGeminiPro), configurado com o modelo "
         "gemini-2.5-flash. Ele recebe o system prompt e o prompt do usuário com as "
         "variáveis mapeadas ({{1.nomeCargo}}, {{1.temaAssunto}}, {{1.tipoTexto}}, "
         "{{1.tomVoz}}) e retorna o texto gerado. O terceiro módulo é o Webhook Response "
         "(gateway:WebhookRespond), que retorna ao frontend a resposta HTTP com status "
         "200 e corpo JSON, mapeando o campo text do Gemini para o campo mensagem da "
         "resposta.")
    body(doc,
         "O blueprint do cenário está versionado no repositório em "
         "make/pulsar-rh-blueprint.json.json, permitindo sua reimportação em qualquer "
         "conta Make.com com um único clique, sem necessidade de reconfigurar os módulos "
         "manualmente.")


def build_s6(doc):
    h1(doc, "6 BENEFÍCIOS PERCEBIDOS E DESAFIOS ENFRENTADOS")

    h2(doc, "6.1 Benefícios Percebidos")
    body(doc,
         "Entre os benefícios observados durante o desenvolvimento e os testes da solução, "
         "destaca-se a redução significativa no tempo de elaboração de comunicações "
         "rotineiras: textos que levariam entre 5 e 15 minutos para serem redigidos "
         "manualmente são gerados em menos de 10 segundos. A qualidade das mensagens "
         "geradas demonstrou-se adequada ao contexto corporativo, mantendo coerência de "
         "tom e formato conforme as opções selecionadas pelo usuário.")
    body(doc,
         "A arquitetura sem servidor próprio (serverless) eliminou custos de infraestrutura "
         "e complexidade operacional: toda a lógica de integração e processamento de IA é "
         "gerenciada externamente, enquanto o frontend é servido gratuitamente pelo GitHub "
         "Pages. Isso torna a solução viável como ferramenta real para pequenas e médias "
         "empresas sem equipe de TI dedicada. Adicionalmente, o blueprint versionado no "
         "repositório garante que qualquer pessoa possa recriar a automação do zero em "
         "minutos.")

    h2(doc, "6.2 Desafios Enfrentados")
    body(doc,
         "O principal desafio técnico foi o refinamento iterativo do prompt. As primeiras "
         "versões do system prompt produziam respostas que incluíam explicações do processo "
         "de geração, listas com marcadores e variações indesejadas de formato. Foram "
         "necessárias múltiplas iterações até chegar à versão atual, que inclui regras "
         "obrigatórias explícitas como 'Não explique o que você está fazendo' e 'Retorne "
         "apenas a mensagem final pronta para uso'.")
    body(doc,
         "Um segundo desafio foi a variabilidade inerente dos LLMs: para entradas "
         "idênticas, o modelo pode gerar saídas ligeiramente diferentes a cada execução. "
         "Embora isso seja geralmente desejável em tarefas criativas, pode gerar "
         "inconsistência em comunicações corporativas onde a padronização é importante. "
         "Por fim, a dependência de serviços externos (Make.com e Google) representa um "
         "ponto único de falha: interrupções nessas plataformas afetam diretamente a "
         "disponibilidade da aplicação, o que exigiria um plano de contingência em um "
         "ambiente de produção corporativo.")


def build_s7(doc):
    h1(doc, "7 LIMITES ÉTICOS E DE SEGURANÇA")

    h2(doc, "7.1 Privacidade de Dados e LGPD")
    body(doc,
         "A aplicação coleta dados pessoais como endereço de e-mail e nome com cargo do "
         "solicitante. De acordo com a Lei Geral de Proteção de Dados Pessoais (Lei "
         "nº 13.709/2018 — LGPD), qualquer tratamento de dados pessoais exige base legal "
         "adequada, que pode ser o legítimo interesse do controlador ou o consentimento "
         "expresso do titular. Em uma implantação corporativa real, seria necessário "
         "formalizar Acordos de Processamento de Dados (DPA) com Make.com e Google, além "
         "de disponibilizar aviso de privacidade acessível ao usuário antes do primeiro uso.")
    body(doc,
         "A versão atual adota o princípio da minimização de dados previsto no artigo 6º "
         "da LGPD: somente os campos estritamente necessários para a geração da mensagem "
         "são coletados, e nenhum dado é armazenado de forma persistente. O system prompt "
         "inclui explicitamente a instrução de não inventar informações, reduzindo o risco "
         "de o modelo gerar dados fictícios sobre pessoas ou situações reais. Os dados "
         "trafegam pela infraestrutura do Make.com (servidor na região us2) e pelo Google, "
         "ambas as plataformas com certificações SOC 2 Type II e conformidade com o GDPR.")

    h2(doc, "7.2 Viés da IA")
    body(doc,
         "LLMs reproduzem padrões presentes em seus dados de treinamento, o que pode "
         "incluir vieses de gênero, hierarquia ou cultura regional. No contexto de "
         "comunicações de RH, esse risco se manifesta de formas específicas: o modelo pode "
         "utilizar formas de tratamento generificadas quando o gênero do destinatário não é "
         "especificado, pode privilegiar um registro linguístico associado a culturas "
         "corporativas específicas de países desenvolvidos, ou pode gerar expressões que "
         "não ressoam igualmente em contextos organizacionais com perfis culturais distintos.")
    body(doc,
         "A mitigação adotada neste projeto inclui a restrição da latitude criativa do "
         "modelo por meio do controle explícito de tom de voz e tipo de texto pelo usuário, "
         "e a instrução de não inventar informações, que ancora a saída no contexto "
         "fornecido. Recomenda-se que, em implantações corporativas, as mensagens geradas "
         "sejam sempre revisadas por um ser humano antes do envio, tratando a IA como "
         "assistente de redação e não como autor final e autônomo.")

    h2(doc, "7.3 Segurança das Credenciais")
    body(doc,
         "A chave de API do Google Gemini está armazenada exclusivamente no cofre seguro "
         "do Make.com, nunca sendo exposta no código-fonte do frontend ou no repositório "
         "público. A URL do webhook do Make.com é injetada em tempo de compilação como "
         "variável de ambiente secreta do GitHub Actions (VITE_MAKE_WEBHOOK_URL), sendo "
         "embarcada no bundle JavaScript sem aparecer no repositório.")
    body(doc,
         "Essa arquitetura segue o princípio do menor privilégio: o frontend possui apenas "
         "permissão para enviar dados a um endpoint público de webhook, sem acesso direto "
         "à API do Gemini nem a qualquer credencial interna. A única superfície de ataque "
         "exposta é a URL do webhook, que pode ser rotacionada no Make.com sem necessidade "
         "de alterar o código-fonte, bastando atualizar o secret no repositório GitHub e "
         "executar um novo deploy automatizado.")


def build_referencias(doc):
    h1(doc, "REFERÊNCIAS")
    refs = [
        "BRASIL. Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de "
        "Dados Pessoais (LGPD). Brasília, DF: Presidência da República, 2018. "
        "Disponível em: https://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm. "
        "Acesso em: abr. 2025.",

        "GOOGLE DEEPMIND. Gemini API Documentation. 2025. Disponível em: "
        "https://ai.google.dev/gemini-api/docs. Acesso em: abr. 2025.",

        "MAKE.COM. Make Help Center: Webhooks. 2025. Disponível em: "
        "https://make.com/en/help. Acesso em: abr. 2025.",

        "META. React: The library for web and native user interfaces. 2025. "
        "Disponível em: https://react.dev. Acesso em: abr. 2025.",

        "OPENAI. OpenAI Platform Documentation. 2025. Disponível em: "
        "https://platform.openai.com/docs. Acesso em: abr. 2025.",

        "ZAPIER. How to use OpenAI with Zapier. 2025. Disponível em: "
        "https://zapier.com/learn/openai. Acesso em: abr. 2025.",

        "WHITE, Jules et al. A Prompt Pattern Catalog to Enhance Prompt Engineering "
        "with ChatGPT. arXiv:2302.11382v1 [cs.SE], 13 fev. 2023. Disponível em: "
        "https://arxiv.org/abs/2302.11382. Acesso em: abr. 2025.",
    ]
    for ref in refs:
        referencia(doc, ref)


# ──────────────────────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────────────────────

def main():
    doc = Document()

    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(3)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    add_page_number(sec)

    build_capa(doc)
    doc.add_page_break()
    build_sumario(doc)
    doc.add_page_break()
    build_s1(doc)
    build_s2(doc)
    build_s3(doc)
    build_s4(doc)
    build_s5(doc)
    build_s6(doc)
    build_s7(doc)
    doc.add_page_break()
    build_referencias(doc)

    doc.save(OUTPUT)
    print(f"Documento salvo em:\n{OUTPUT}")


if __name__ == "__main__":
    main()
